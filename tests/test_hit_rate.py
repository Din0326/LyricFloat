"""Hit-rate tool logic. Songs and lyric lines are made up; sources are fakes."""

import csv

import pytest

from lyricfloat.diagnostics.hit_rate import (
    ListedSong,
    adopt_from_xlsx,
    build_candidates,
    count_simplified,
    evaluate,
    load_test_list,
    summarize,
    write_candidates_xlsx,
    write_report_xlsx,
)
from lyricfloat.diagnostics.problem_report import add_to_test_list
from lyricfloat.lyrics.chinese import ChineseConverter
from lyricfloat.lyrics.pipeline import OpenCCPipeline
from lyricfloat.lyrics.service import LyricResolver
from lyricfloat.lyrics.sources.base import LyricSource
from lyricfloat.models import Candidate, LyricLine, LyricResult, Lyrics, ResultKind, TrackInfo

PIPE = OpenCCPipeline(converter=ChineseConverter())


def history_item(title, artist="範例樂團", at="2026-10-07T12:00:00", kind="synced", language="zh"):
    return {"at": at, "title": title, "artist": artist, "album": "", "duration": 200.0, "source_app": "Spotify.exe",
            "language": language, "kind": kind, "source": "lrclib", "score": 0.9}


def test_load_test_list_reads_old_and_new_headers(tmp_path):
    old = tmp_path / "old.csv"
    old.write_text("title,artist,album,duration,language,source_app,added,note\n紙船,範例樂團,,200,zh,Spotify.exe,2026-10-07,\n", encoding="utf-8-sig")
    (song,) = load_test_list(old)
    assert song.title == "紙船" and song.duration == 200.0 and song.category == ""
    assert load_test_list(tmp_path / "missing.csv") == []


def test_adding_to_an_old_file_keeps_its_columns(tmp_path):
    old = tmp_path / "old.csv"
    old.write_text("title,artist,album,duration,language,source_app,added,note\n", encoding="utf-8-sig")
    assert add_to_test_list(TrackInfo("Paper Boats", "The Example Band", "", 180.0), "en", "", old, category="冷門")
    with old.open(newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    assert rows[0]["title"] == "Paper Boats" and "category" not in rows[0]


def test_candidates_dedupe_count_and_flag_listed():
    history = [
        history_item("紙船", at="2026-10-07T10:00:00"),
        history_item("紙船", at="2026-10-07T11:00:00", kind="not_found"),
        history_item("Paper Boats", "The Example Band", language="en"),
    ]
    listed = [ListedSong("Paper Boats", "The Example Band")]
    rows = build_candidates(history, listed)
    assert [(r["title"], r["plays"], r["listed"]) for r in rows] == [("紙船", 2, "否"), ("Paper Boats", 1, "是")]
    assert rows[0]["last_kind"] == "not_found" and rows[0]["last_played"] == "2026-10-07T11:00:00"


def test_candidates_workbook_round_trip_through_adopt(tmp_path):
    from openpyxl import load_workbook

    rows = build_candidates([history_item("紙船"), history_item("Paper Boats", "The Example Band", language="en")], [])
    xlsx = tmp_path / "candidates.xlsx"
    write_candidates_xlsx(xlsx, rows, [])
    wb = load_workbook(xlsx)
    ws = wb["候選"]
    header = [c.value for c in ws[1]]
    assert header[:3] == ["納入", "語言", "類型"]
    assert ws.freeze_panes == "D2" and ws.auto_filter.ref
    assert ws.data_validations.dataValidation  # drop-downs on the fill-in columns
    ws.cell(row=2, column=1, value="是")
    ws.cell(row=2, column=3, value="冷門")
    wb.save(xlsx)
    summary = load_workbook(xlsx)["摘要"]
    assert str(summary["B2"].value).startswith("=COUNTIFS(")

    target = tmp_path / "test_list.csv"
    assert adopt_from_xlsx(xlsx, target) == (1, 0)
    assert adopt_from_xlsx(xlsx, target) == (0, 1)
    (song,) = load_test_list(target)
    assert song.title == rows[0]["title"] and song.category == "冷門" and song.language == rows[0]["language"]


def test_count_simplified_covers_everything_displayed():
    lyrics = Lyrics((LyricLine(1, "这里是台湾"), LyricLine(2, "台灣的歌"), LyricLine(3, "私の国")), True, ("软件", None, None))
    result = LyricResult(TrackInfo("x", "y"), ResultKind.SYNCED, lyrics)
    shown = TrackInfo("请让一下", "範例樂團")
    # simplified: one lyric line, one translation, the title; Japanese and traditional lines are not counted
    assert count_simplified(result, shown, PIPE.cc.is_simplified) == 3


# --- evaluation with fake sources ------------------------------------------------------


class Fake(LyricSource):
    name = "lrclib"

    def __init__(self, found):
        self.found = found

    async def search(self, track):
        return [c for c in self.found if c.title == track.title]


def cand(title, lrc="[00:01.00]made up\n[00:05.00]also made up"):
    return Candidate("lrclib", title, title, "範例樂團", "", 200.0, lrc, None)


async def test_evaluate_and_summarize():
    songs = [
        ListedSong("Found One", "範例樂團", duration=200.0, language="en"),
        ListedSong("Missing", "範例樂團", duration=200.0, language="en"),
        ListedSong("紙船", "範例樂團", duration=200.0, language="zh"),
        ListedSong("简体歌", "範例樂團", duration=200.0, language="zh"),
    ]
    found = [cand("Found One"), cand("紙船"), cand("简体歌", "[00:01.00]软件更新")]
    resolver = LyricResolver([Fake(found)], pipeline=PIPE)
    seen = []
    outcomes = await evaluate(songs, resolver, PIPE, delay=0, progress=lambda i, n, o: seen.append(i))
    assert seen == [1, 2, 3, 4]
    assert [o.synced for o in outcomes] == [True, False, True, True]
    assert outcomes[1].kind == "not_found"
    assert outcomes[3].simplified == 0  # lyrics converted
    assert outcomes[3].first_ms is not None
    s = summarize(outcomes)
    assert s["languages"]["en"]["rate"] == pytest.approx(0.5) and s["languages"]["en"]["passed"] is False
    assert s["languages"]["zh"]["rate"] == pytest.approx(1.0) and s["languages"]["zh"]["passed"] is True
    assert s["languages"]["other"]["passed"] is None
    assert not s["passed"]


async def test_report_workbook_uses_formulas(tmp_path):
    from openpyxl import load_workbook

    resolver = LyricResolver([Fake([cand("Found One")])], pipeline=PIPE)
    outcomes = await evaluate([ListedSong("Found One", "範例樂團", duration=200.0, language="en")], resolver, PIPE, delay=0)
    path = tmp_path / "report.xlsx"
    write_report_xlsx(path, outcomes)
    wb = load_workbook(path)
    assert wb.sheetnames == ["摘要", "明細", "說明"]
    formulas = [c.value for c in wb["摘要"]["B"] if isinstance(c.value, str) and c.value.startswith("=")]
    assert any(f.startswith("=COUNTIFS(明細!") for f in formulas)
    assert any("通過" in f for f in formulas)
    detail = wb["明細"]
    assert detail["H2"].value == "是"  # synced hit
