from types import SimpleNamespace

import pytest

from lyricfloat.diagnostics.perf import (
    PerfSample,
    find_app_process,
    read_samples,
    summarize,
    take_sample,
    verdict,
    write_perf_xlsx,
    write_samples,
)


def samples(rss_values, cpu=0.3, alive_last=True):
    out = [PerfSample(1000.0 + i * 60, i * 60.0, cpu, cpu * 16, rss * 2, rss * 1.2, pws_mb=rss) for i, rss in enumerate(rss_values)]
    if not alive_last:
        out.append(PerfSample(0, out[-1].elapsed + 60, 0, 0, 0, 0, alive=False))
    return out


def test_idle_run_passes_cpu_and_memory():
    s = summarize("idle", samples([100.0] * 20, cpu=0.2))
    assert s["cpu_median"] == pytest.approx(0.2) and s["cpu_core_median"] == pytest.approx(3.2)
    assert verdict(s) == {"memory": True, "cpu": True}


def test_playing_uses_its_own_cpu_target():
    assert verdict(summarize("playing", samples([100.0] * 5, cpu=2.5)))["cpu"] is True
    assert verdict(summarize("idle", samples([100.0] * 5, cpu=2.5)))["cpu"] is False


def test_memory_verdict_uses_the_median_private_working_set():
    # A brief conversion peak does not fail the run; a high steady level does.
    assert verdict(summarize("idle", samples([60.0, 146.0, 62.0])))["memory"] is True
    assert verdict(summarize("idle", samples([125.0, 130.0, 128.0])))["memory"] is False


def test_soak_growth_compares_first_and_last_tenth():
    rss = [100.0] * 10 + [105.0] * 80 + [108.0] * 10
    s = summarize("soak", samples(rss))
    assert s["growth_pct"] == pytest.approx(8.0)
    assert verdict(s) == {"memory": True, "no_crash": True, "growth": True}


def test_soak_crash_is_reported():
    s = summarize("soak", samples([100.0] * 12, alive_last=False))
    assert s["crashed"] and verdict(s)["no_crash"] is False


def test_csv_round_trip_and_append(tmp_path):
    path = tmp_path / "perf.csv"
    write_samples(path, samples([100.0, 101.0]), {"label": "soak"}, append=True)
    write_samples(path, samples([102.0]), {"label": "soak"}, append=True)
    meta, loaded = read_samples(path)
    assert meta == {"label": "soak"}
    assert [s.pws_mb for s in loaded] == [100.0, 101.0, 102.0]  # the judged figure survives the CSV
    assert [s.rss_mb for s in loaded] == [200.0, 202.0, 204.0]


def test_take_sample_normalises_cpu_to_the_whole_machine():
    proc = SimpleNamespace(
        cpu_percent=lambda interval=None: 8.0,
        memory_info=lambda: SimpleNamespace(rss=110 * 1048576, private=70 * 1048576),
        memory_full_info=lambda: SimpleNamespace(uss=62 * 1048576),
    )
    s = take_sample(proc, started=100.0, now_wall=160.0, cpu_count=16)
    assert s.cpu_pct == pytest.approx(0.5) and s.cpu_core_pct == pytest.approx(8.0)
    assert s.rss_mb == pytest.approx(110.0) and s.private_mb == pytest.approx(70.0) and s.elapsed == 60.0
    assert s.pws_mb == pytest.approx(62.0)


def test_find_app_process_picks_the_interpreter_not_the_launcher():
    def proc(cmd, rss, name="python.exe"):
        return SimpleNamespace(info={"name": name, "cmdline": cmd.split(), "memory_info": SimpleNamespace(rss=rss)})

    launcher = proc(r"C:\x\.venv\Scripts\python.exe -m lyricfloat", 7_000_000)
    app = proc(r"C:\Python312\python.exe -m lyricfloat", 110_000_000)
    other = proc("python.exe -m something_else", 500_000_000)
    fake_psutil = SimpleNamespace(
        process_iter=lambda attrs: [launcher, app, other], NoSuchProcess=Exception, AccessDenied=Exception
    )
    assert find_app_process(fake_psutil) is app
    fake_psutil.process_iter = lambda attrs: [other]
    assert find_app_process(fake_psutil) is None


def test_summary_workbook(tmp_path):
    from openpyxl import load_workbook

    run = summarize("idle", samples([100.0] * 5))
    run["started"] = "2026-10-07 13:30:00"
    path = tmp_path / "perf.xlsx"
    write_perf_xlsx(path, [run])
    ws = load_workbook(path)["摘要"]
    assert ws["A2"].value == "idle"
    assert "記憶體通過" in ws.cell(row=2, column=ws.max_column).value
