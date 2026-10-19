"""L3 statistics workbook (numbers only)."""

import json

from openpyxl import load_workbook

from lyricfloat.diagnostics.alignment_report import alignment_log_files, read_records, write_alignment_xlsx

RECORDS = [
    {"at": "2026-10-07 18:00:00", "language": "zh", "ok": True, "usable": True, "lines": 40, "judged": 38, "sure": 30,
     "sure_share": 0.79, "median_abs_diff_s": 0.08, "within_0.3s": 0.8, "median_shift_s": 0.05, "mean_confidence": 0.41,
     "word_lines": 28, "seconds": 220.0, "separate_s": 50.1, "align_s": 2.0, "wall_s": 60.2},
    {"at": "2026-10-07 18:05:00", "language": "en", "ok": False, "seconds": 190.0, "wall_s": 3.0},
    {"at": "2026-10-07 18:09:00", "event": "dropped", "reason": "paused"},
]


def test_log_files_are_read_oldest_first(tmp_path):
    (tmp_path / "alignment.jsonl.1").write_text(json.dumps({"at": "old"}) + "\n", encoding="utf-8")
    (tmp_path / "alignment.jsonl").write_text(json.dumps({"at": "new"}) + "\nnot json\n", encoding="utf-8")
    assert [r["at"] for r in read_records(alignment_log_files(tmp_path))] == ["old", "new"]


def test_workbook_has_rows_and_formula_summary(tmp_path):
    out = tmp_path / "report.xlsx"
    write_alignment_xlsx(out, RECORDS)
    wb = load_workbook(out)
    assert wb.sheetnames == ["摘要", "明細", "沒錄成", "說明"]
    detail = wb["明細"]
    assert detail.max_row == 3 and detail["B2"].value == "zh" and detail["D2"].value == "True"
    assert detail.freeze_panes == "C2" and detail.auto_filter.ref.startswith("A1:")
    formulas = [c.value for c in wb["摘要"]["B"] if c.value is not None]
    assert all(isinstance(v, str) and v.startswith("=") for v in formulas)


def test_dropped_recordings_get_their_own_sheet(tmp_path):
    out = tmp_path / "report.xlsx"
    write_alignment_xlsx(out, RECORDS)
    wb = load_workbook(out)
    assert wb["明細"].max_row == 3  # dropped plays are not aligned songs
    assert wb["沒錄成"]["B2"].value == "中途暫停"


def test_summary_formulas_point_at_the_right_columns(tmp_path):
    out = tmp_path / "report.xlsx"
    write_alignment_xlsx(out, RECORDS)
    wb = load_workbook(out)
    headers = {c.value: c.column_letter for c in wb["明細"][1]}
    summary = {row[0].value: row[1].value for row in wb["摘要"].iter_rows() if row[0].value}
    assert f"明細!{headers['整體 (s)']}2:" in summary["整體耗時中位數 (s)"]
    assert f"明細!{headers['0.3 秒內比例']}2:" in summary["華語：0.3 秒內比例（平均）"]
