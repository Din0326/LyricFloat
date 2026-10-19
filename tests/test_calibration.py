from lyricfloat.diagnostics.calibration import (
    CalibrationRecorder,
    nearest_line,
    read_marks,
    summarize_marks,
)


def test_nearest_line_within_three_seconds():
    starts = [5.0, 10.0, 15.0]
    assert nearest_line(starts, 10.4) == 1
    assert nearest_line(starts, 12.4) == 1
    assert nearest_line(starts, 12.6) == 2
    assert nearest_line(starts, 2.1) == 0  # 2.9 s before the first line
    assert nearest_line(starts, 1.9) is None  # more than 3 s before the first line
    assert nearest_line([], 3.0) is None


def test_recorder_counts_survive_a_restart(tmp_path):
    path = tmp_path / "calibration.jsonl"
    rec = CalibrationRecorder(path)
    assert rec.record("紙船", "範例樂團", "lrclib", "1", [5.0, 10.0], 0.0, 5.3).delta_ms == 300.0
    assert CalibrationRecorder(path).count("紙船", "範例樂團") == 1


def test_summary_median_of_absolute_errors(tmp_path):
    path = tmp_path / "calibration.jsonl"
    rec = CalibrationRecorder(path)
    for pos in (5.2, 4.7, 5.9):  # +200, -300, +900 ms
        rec.record("紙船", "範例樂團", "lrclib", "1", [5.0], 0.0, pos)
    s = summarize_marks(read_marks(path))
    assert s["marks"] == 3 and s["songs"] == 1
    assert s["median_abs_ms"] == 300.0 and s["median_signed_ms"] == 200.0
    assert s["passed"] is True
    assert summarize_marks([])["passed"] is None


def test_marks_workbook_uses_formulas(tmp_path):
    from openpyxl import load_workbook

    from lyricfloat.diagnostics.calibration import write_marks_xlsx

    path = tmp_path / "calibration.jsonl"
    rec = CalibrationRecorder(path)
    rec.record("紙船", "範例樂團", "lrclib", "1", [5.0, 10.0], 0.0, 5.3)
    out = tmp_path / "marks.xlsx"
    write_marks_xlsx(out, read_marks(path))
    wb = load_workbook(out)
    assert wb.sheetnames == ["摘要", "標記", "說明"]
    assert wb["標記"]["K2"].value == "=ABS(J2)"
    assert str(wb["摘要"]["B3"].value).startswith('=IF(B1=0,"—",MEDIAN(')
