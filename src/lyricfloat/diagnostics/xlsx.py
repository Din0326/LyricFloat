"""Workbook helpers following Din's table conventions.

Microsoft JhengHei, header row fill, thin borders, frozen header with
filter, wrapped long text with row heights, fill-in columns in yellow
with drop-downs, landscape fit-to-width printing. Statistics are written
as formulas by the callers (COUNTIFS), never as fixed numbers.
"""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass, field
from typing import Any

FONT = "微軟正黑體"
HEADER_FILL = "DDE7F0"
INPUT_FILL = "FFF9E0"
BORDER_COLOR = "BFBFBF"


@dataclass(frozen=True)
class Column:
    key: str
    title: str
    width: float = 12
    wrap: bool = False
    fill_in: bool = False  # yellow: for the user to fill
    choices: Sequence[str] = field(default_factory=tuple)  # drop-down values
    colors: dict[str, str] = field(default_factory=dict)  # value -> fill colour


def _units(text: str) -> int:
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in text)


def write_table(ws, columns: Sequence[Column], rows: Sequence[dict[str, Any]], freeze: str = "A2") -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    thin = Side(style="thin", color=BORDER_COLOR)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c, col in enumerate(columns, 1):
        cell = ws.cell(row=1, column=c, value=col.title)
        cell.font = Font(name=FONT, bold=True, size=11)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(c)].width = col.width
    ws.row_dimensions[1].height = 32

    last = max(len(rows) + 1, 2)
    for r, row in enumerate(rows, 2):
        lines = 1
        for c, col in enumerate(columns, 1):
            value = row.get(col.key)
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=col.wrap, horizontal="left" if col.wrap else "center")
            cell.border = border
            if col.fill_in:
                cell.fill = PatternFill("solid", fgColor=INPUT_FILL)
            elif value is not None and str(value) in col.colors:
                cell.fill = PatternFill("solid", fgColor=col.colors[str(value)])
            if col.wrap and value:
                lines = max(lines, -(-_units(str(value)) // max(1, int(col.width))))
        ws.row_dimensions[r].height = max(18, 15 * lines)

    for c, col in enumerate(columns, 1):
        if col.choices:
            letter = get_column_letter(c)
            dv = DataValidation(type="list", formula1='"' + ",".join(col.choices) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{max(last, 500)}")

    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def write_stats(ws, start_row: int, stats: Sequence[tuple[str, Any]], label_width: float = 34) -> None:
    """Label / value pairs in columns A and B; a value is a number or a formula ("=COUNTIFS(...)")."""
    from openpyxl.styles import Alignment, Font

    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, label_width)
    for i, (label, formula) in enumerate(stats):
        row = start_row + i
        a = ws.cell(row=row, column=1, value=label)
        a.font = Font(name=FONT, size=10, bold=True)
        b = ws.cell(row=row, column=2, value=formula)
        b.font = Font(name=FONT, size=10)
        b.alignment = Alignment(horizontal="center")


def write_text_sheet(ws, paragraphs: Sequence[str], width: float = 110) -> None:
    from openpyxl.styles import Alignment, Font

    ws.column_dimensions["A"].width = width
    for r, para in enumerate(paragraphs, 1):
        heading = para.startswith("#")
        cell = ws.cell(row=r, column=1, value=para.lstrip("# ") if heading else para)
        cell.font = Font(name=FONT, size=11, bold=heading)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = max(18, 16 * -(-_units(para) // int(width - 5)))
