"""Analysis of SMTC timeline samples for the M1 gate (player progress measurement).

Reference: during uninterrupted playback the true position advances 1 s
per wall-clock second. Within each playing segment every position series
is fitted to pos = elapsed + c (c = that series' median offset, so a
constant bias is not measurable here) and errors are the residuals:

- naive:   the raw Position as reported, treated as "now"
- extrap:  Position + (now - LastUpdatedTime), what MediaWatcher computes
- engine:  what SyncEngine shows, anchored only when a new report arrives

Absolute offset between reported position and audible audio is out of
scope (that needs audio capture, see L2).
"""

from __future__ import annotations

import csv
import statistics
from collections.abc import Iterable, Sequence
from dataclasses import asdict, dataclass, field, fields
from pathlib import Path

from lyricfloat.media.timeline import RawTimeline, anchor_from_timeline
from lyricfloat.models import PlaybackStatus
from lyricfloat.sync.engine import SEEK_THRESHOLD, SyncEngine

MIN_SEGMENT_SECONDS = 15.0
BURST_GAP = 0.25  # reports closer than this belong to the same state change
GATE_CORRECTION_P95_MS = 100.0  # M1 gate (SPEC v1.5): correction size p95 <= 100 ms


@dataclass(slots=True)
class Sample:
    t_mono: float
    t_wall: float
    kind: str  # "poll", "event:timeline", "event:playback", "event:media", "action:<name>"
    status: int
    position: float
    duration: float
    last_updated: float
    title: str = ""
    detail: str = ""

    @property
    def raw(self) -> RawTimeline:
        return RawTimeline(self.position, self.duration, self.last_updated, PlaybackStatus(self.status))

    def extrapolated(self) -> float:
        return anchor_from_timeline(self.raw, self.t_wall, self.t_mono).position


SAMPLE_FIELDS = [f.name for f in fields(Sample)]


def write_samples(path: Path, samples: Iterable[Sample], meta: dict[str, str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        for key, value in meta.items():
            fh.write(f"# {key}={value}\n")
        writer = csv.DictWriter(fh, fieldnames=SAMPLE_FIELDS)
        writer.writeheader()
        for s in samples:
            writer.writerow(asdict(s))


def read_samples(path: Path) -> tuple[dict[str, str], list[Sample]]:
    meta: dict[str, str] = {}
    rows: list[str] = []
    with path.open(encoding="utf-8") as fh:
        for line in fh:
            if line.startswith("# ") and "=" in line:
                key, value = line[2:].rstrip("\n").split("=", 1)
                meta[key] = value
            else:
                rows.append(line)
    samples = []
    for row in csv.DictReader(rows):
        samples.append(
            Sample(
                t_mono=float(row["t_mono"]),
                t_wall=float(row["t_wall"]),
                kind=row["kind"],
                status=int(row["status"]),
                position=float(row["position"]),
                duration=float(row["duration"]),
                last_updated=float(row["last_updated"]),
                title=row.get("title", ""),
                detail=row.get("detail", ""),
            )
        )
    return meta, samples


def _pct(values: Sequence[float], q: float) -> float | None:
    if not values:
        return None
    ordered = sorted(values)
    k = (len(ordered) - 1) * q
    lo, hi = int(k), min(int(k) + 1, len(ordered) - 1)
    return ordered[lo] + (ordered[hi] - ordered[lo]) * (k - lo)


def _median(values: Sequence[float]) -> float | None:
    return statistics.median(values) if values else None


def _slope(xs: Sequence[float], ys: Sequence[float]) -> float | None:
    if len(xs) < 3:
        return None
    mx, my = statistics.fmean(xs), statistics.fmean(ys)
    var = sum((x - mx) ** 2 for x in xs)
    if var == 0:
        return None
    return sum((x - mx) * (y - my) for x, y in zip(xs, ys)) / var


@dataclass(slots=True)
class Segment:
    index: int
    start: float  # seconds since the run started
    length: float
    samples: int
    updates: int
    end_reason: str
    naive_abs_ms: list[float] = field(default_factory=list)
    extrap_abs_ms: list[float] = field(default_factory=list)
    engine_abs_ms: list[float] = field(default_factory=list)
    correction_ms: list[float] = field(default_factory=list)
    drift_ms_per_min: float | None = None

    @property
    def usable(self) -> bool:
        return self.length >= MIN_SEGMENT_SECONDS


@dataclass(slots=True)
class RunSummary:
    player: str
    file: str
    app_id: str
    started: str
    seconds: float
    playing_seconds: float
    tracks: int
    segments: list[Segment]
    update_intervals: list[float]
    timeline_events: int
    playback_events: int
    media_events: int
    pauses: int
    seeks: int
    action_latency_ms: dict[str, list[float]]

    def _pooled(self, attr: str) -> list[float]:
        return [v for s in self.segments if s.usable for v in getattr(s, attr)]

    def stat(self, attr: str, q: float) -> float | None:
        return _pct(self._pooled(attr), q)

    @property
    def drift(self) -> float | None:
        return _median([s.drift_ms_per_min for s in self.segments if s.usable and s.drift_ms_per_min is not None])

    @property
    def timeline_events_per_min(self) -> float | None:
        return self.timeline_events / (self.playing_seconds / 60) if self.playing_seconds > 0 else None

    @property
    def updates(self) -> int:
        return sum(s.updates for s in self.segments)

    @property
    def updates_per_min(self) -> float | None:
        return self.updates / (self.playing_seconds / 60) if self.playing_seconds > 0 else None

    @property
    def corrections(self) -> int:
        return len(self._pooled("correction_ms"))


def _residuals_ms(series: Sequence[float], elapsed: Sequence[float]) -> list[float]:
    """|series - (elapsed + c)| with c fitted to this series alone (median offset).

    Each method is judged on its own consistency against a 1x clock; a
    constant offset is not measurable without audio, so it is not counted.
    """
    c = statistics.median(p - e for p, e in zip(series, elapsed))
    return [abs(p - (e + c)) * 1000 for p, e in zip(series, elapsed)]


def _analyse_segment(index: int, run_start: float, polls: list[Sample], update_times: list[float], end_reason: str) -> Segment:
    t0 = polls[0].t_mono
    elapsed = [s.t_mono - t0 for s in polls]
    naive = [s.position for s in polls]
    extrap = [s.extrapolated() for s in polls]

    # SyncEngine anchored only on samples that carry a new report. Each
    # re-anchor shows how far the displayed position had wandered from what
    # the player now reports: the correction size.
    engine = SyncEngine(clock=lambda: 0.0)
    shown: list[float] = []
    corrections: list[float] = []
    last_sig = None
    for s in polls:
        if s.raw.signature != last_sig:
            last_sig = s.raw.signature
            anchor = anchor_from_timeline(s.raw, s.t_wall, s.t_mono)
            predicted = engine.position(s.t_mono)
            if predicted is not None:
                corrections.append(abs(anchor.position - predicted) * 1000)
            engine.update(anchor)
        shown.append(engine.position(s.t_mono))

    slope = _slope(elapsed, extrap)
    return Segment(
        index=index,
        start=t0 - run_start,
        length=elapsed[-1],
        samples=len(polls),
        updates=len(update_times),
        end_reason=end_reason,
        naive_abs_ms=_residuals_ms(naive, elapsed),
        extrap_abs_ms=_residuals_ms(extrap, elapsed),
        engine_abs_ms=_residuals_ms(shown, elapsed),
        correction_ms=corrections,
        drift_ms_per_min=None if slope is None else (slope - 1.0) * 60_000,
    )


def analyse(player: str, file: str, meta: dict[str, str], samples: list[Sample]) -> RunSummary:
    polls = [s for s in samples if s.kind == "poll"]
    run_start = samples[0].t_mono if samples else 0.0
    segments: list[Segment] = []
    intervals: list[float] = []
    pauses = seeks = 0
    titles: set[str] = set()

    current: list[Sample] = []
    updates: list[float] = []
    last_sig = None
    prev: Sample | None = None

    def close(reason: str) -> None:
        nonlocal current, updates
        if len(current) >= 3:
            segments.append(_analyse_segment(len(segments) + 1, run_start, current, updates, reason))
        current, updates = [], []

    for s in polls:
        if s.title:
            titles.add(s.title)
        playing = s.status == PlaybackStatus.PLAYING
        if prev is not None:
            if prev.status == PlaybackStatus.PLAYING and s.status == PlaybackStatus.PAUSED:
                pauses += 1
            if s.title != prev.title:
                close("track")
            elif not playing:
                if current:
                    close("pause")
            elif current:
                predicted = current[-1].extrapolated() + (s.t_mono - current[-1].t_mono)
                if abs(s.extrapolated() - predicted) > SEEK_THRESHOLD:
                    seeks += 1
                    close("seek")
        if playing:
            sig = s.raw.signature
            if sig != last_sig:
                last_sig = sig
                # Several reports for one state change count as one update.
                if not updates or s.t_mono - updates[-1] >= BURST_GAP:
                    if current and updates:
                        intervals.append(s.t_mono - updates[-1])
                    updates.append(s.t_mono)
            current.append(s)
        else:
            last_sig = None
        prev = s
    close("end")

    playing_seconds = sum(
        b.t_mono - a.t_mono for a, b in zip(polls, polls[1:]) if a.status == PlaybackStatus.PLAYING
    )
    latency = _action_latency(samples)
    return RunSummary(
        player=player,
        file=file,
        app_id=meta.get("app_id", ""),
        started=meta.get("started", ""),
        seconds=(polls[-1].t_mono - polls[0].t_mono) if len(polls) > 1 else 0.0,
        playing_seconds=playing_seconds,
        tracks=len(titles),
        segments=segments,
        update_intervals=intervals,
        timeline_events=sum(1 for s in samples if s.kind == "event:timeline"),
        playback_events=sum(1 for s in samples if s.kind == "event:playback"),
        media_events=sum(1 for s in samples if s.kind == "event:media"),
        pauses=pauses,
        seeks=seeks,
        action_latency_ms=latency,
    )


def _action_latency(samples: list[Sample]) -> dict[str, list[float]]:
    """Time from a scripted SMTC command to the first poll that reflects it."""
    out: dict[str, list[float]] = {}
    polls = [s for s in samples if s.kind == "poll"]
    for action in (s for s in samples if s.kind.startswith("action:")):
        name = action.kind.split(":", 1)[1]
        target = float(action.detail) if action.detail else None
        for s in polls:
            if s.t_mono <= action.t_mono:
                continue
            hit = (
                (name == "pause" and s.status == PlaybackStatus.PAUSED)
                or (name == "play" and s.status == PlaybackStatus.PLAYING)
                or (name == "seek" and target is not None and abs(s.extrapolated() - target) < 1.5)
            )
            if hit:
                out.setdefault(name, []).append((s.t_mono - action.t_mono) * 1000)
                break
            if s.t_mono - action.t_mono > 10:
                out.setdefault(name + "_timeout", []).append(10_000.0)
                break
    return out


# ---------------------------------------------------------------- reporting


def _r(value: float | None, digits: int = 0) -> float | None:
    return None if value is None else round(value, digits)


def gate_verdict(corr_p95: float | None) -> str:
    if corr_p95 is None:
        return "無資料"
    return "通過" if corr_p95 <= GATE_CORRECTION_P95_MS else "未通過"


def summary_row(run: RunSummary) -> dict[str, object]:
    lat = run.action_latency_ms
    corr_p95 = run.stat("correction_ms", 0.95)
    return {
        "player": run.player,
        "status": "已量測",
        "note": "",
        "started": run.started,
        "app_id": run.app_id,
        "seconds": _r(run.seconds, 1),
        "playing": _r(run.playing_seconds, 1),
        "tracks": run.tracks,
        "segments": sum(1 for s in run.segments if s.usable),
        "updates": run.updates,
        "updates_min": _r(run.updates_per_min, 1),
        "interval_med": _r(_median(run.update_intervals), 2),
        "interval_p95": _r(_pct(run.update_intervals, 0.95), 2),
        "tl_events_min": _r(run.timeline_events_per_min, 1),
        "naive_p95": _r(run.stat("naive_abs_ms", 0.95)),
        "naive_max": _r(run.stat("naive_abs_ms", 1.0)),
        "extrap_med": _r(run.stat("extrap_abs_ms", 0.5)),
        "extrap_p95": _r(run.stat("extrap_abs_ms", 0.95)),
        "engine_med": _r(run.stat("engine_abs_ms", 0.5)),
        "engine_p95": _r(run.stat("engine_abs_ms", 0.95)),
        "engine_max": _r(run.stat("engine_abs_ms", 1.0)),
        "corrections": run.corrections,
        "corr_med": _r(run.stat("correction_ms", 0.5)),
        "corr_p95": _r(corr_p95),
        "corr_max": _r(run.stat("correction_ms", 1.0)),
        "gate": gate_verdict(corr_p95),
        "drift": _r(run.drift, 1),
        "pauses": run.pauses,
        "seeks": run.seeks,
        "lat_pause": _r(_median(lat.get("pause", [])), 0),
        "lat_play": _r(_median(lat.get("play", [])), 0),
        "lat_seek": _r(_median(lat.get("seek", [])), 0),
        "file": run.file,
    }


SUMMARY_COLUMNS = [
    ("player", "播放器", 12, False),
    ("status", "狀態", 10, False),
    ("gate", "M1 閘門判定", 12, False),
    ("note", "說明", 30, True),
    ("started", "量測時間", 18, False),
    ("app_id", "SMTC 來源 ID", 18, True),
    ("seconds", "取樣秒數", 10, False),
    ("playing", "播放中秒數", 10, False),
    ("tracks", "歌曲數", 8, False),
    ("segments", "有效區段數", 10, False),
    ("updates", "進度更新次數", 12, False),
    ("updates_min", "更新次數／分", 10, False),
    ("interval_med", "更新間隔中位數 (s)", 12, False),
    ("interval_p95", "更新間隔 p95 (s)", 12, False),
    ("tl_events_min", "timeline 事件／分", 12, False),
    ("naive_p95", "直接取值誤差 p95 (ms)", 14, False),
    ("naive_max", "直接取值誤差最大 (ms)", 14, False),
    ("extrap_med", "外推誤差中位數 (ms)", 14, False),
    ("extrap_p95", "外推誤差 p95 (ms)", 14, False),
    ("engine_med", "SyncEngine 誤差中位數 (ms)", 16, False),
    ("engine_p95", "SyncEngine 誤差 p95 (ms)", 16, False),
    ("engine_max", "SyncEngine 誤差最大 (ms)", 16, False),
    ("corrections", "對時次數", 10, False),
    ("corr_med", "對時校正量中位數 (ms)", 14, False),
    ("corr_p95", "對時校正量 p95 (ms)", 14, False),
    ("corr_max", "對時校正量最大 (ms)", 14, False),
    ("drift", "漂移 (ms／分)", 12, False),
    ("pauses", "偵測暫停次數", 10, False),
    ("seeks", "偵測拖動次數", 10, False),
    ("lat_pause", "暫停回報延遲 (ms)", 12, False),
    ("lat_play", "播放回報延遲 (ms)", 12, False),
    ("lat_seek", "拖動回報延遲 (ms)", 12, False),
    ("file", "原始資料檔", 34, True),
]

SEGMENT_COLUMNS = [
    ("player", "播放器", 12, False),
    ("started", "量測時間", 18, False),
    ("index", "區段", 8, False),
    ("start", "開始 (s)", 10, False),
    ("length", "長度 (s)", 10, False),
    ("samples", "樣本數", 10, False),
    ("updates", "更新次數", 10, False),
    ("end_reason", "結束原因", 10, False),
    ("usable", "納入統計", 10, False),
    ("naive_p95", "直接取值 p95 (ms)", 14, False),
    ("extrap_p95", "外推 p95 (ms)", 14, False),
    ("engine_p95", "SyncEngine p95 (ms)", 14, False),
    ("engine_max", "SyncEngine 最大 (ms)", 14, False),
    ("drift", "漂移 (ms／分)", 12, False),
]

END_REASON_TEXT = {"track": "切歌", "pause": "暫停", "seek": "拖動", "end": "量測結束"}


def segment_rows(run: RunSummary) -> list[dict[str, object]]:
    return [
        {
            "player": run.player,
            "started": run.started,
            "index": s.index,
            "start": _r(s.start, 1),
            "length": _r(s.length, 1),
            "samples": s.samples,
            "updates": s.updates,
            "end_reason": END_REASON_TEXT.get(s.end_reason, s.end_reason),
            "usable": "是" if s.usable else "否（< 15 秒）",
            "naive_p95": _r(_pct(s.naive_abs_ms, 0.95)),
            "extrap_p95": _r(_pct(s.extrap_abs_ms, 0.95)),
            "engine_p95": _r(_pct(s.engine_abs_ms, 0.95)),
            "engine_max": _r(_pct(s.engine_abs_ms, 1.0)),
            "drift": _r(s.drift_ms_per_min, 1),
        }
        for s in run.segments
    ]


def text_summary(rows: list[dict[str, object]]) -> str:
    lines = []
    for row in rows:
        r = {k: ("—" if v is None else v) for k, v in row.items()}
        if r["status"] != "已量測":
            lines.append(f"[{r['player']}] {r['status']}：{r['note']}")
            continue
        lines.append(
            f"[{r['player']}] {r['started']} 播放 {r['playing']}s／更新 {r['updates']} 次（{r['updates_min']}／分），"
            f"間隔中位數 {r['interval_med']}s（p95 {r['interval_p95']}s），timeline 事件 {r['tl_events_min']}／分\n"
            f"    直接取值誤差 p95 {r['naive_p95']}ms（最大 {r['naive_max']}ms）｜外推 p95 {r['extrap_p95']}ms｜"
            f"SyncEngine 中位數 {r['engine_med']}ms、p95 {r['engine_p95']}ms｜漂移 {r['drift']}ms／分\n"
            f"    對時 {r['corrections']} 次，校正量中位數 {r['corr_med']}ms、p95 {r['corr_p95']}ms、最大 {r['corr_max']}ms"
            f"｜M1 閘門：{r['gate']}\n"
            f"    暫停 {r['pauses']} 次、拖動 {r['seeks']} 次｜回報延遲 暫停 {r['lat_pause']}ms、播放 {r['lat_play']}ms、拖動 {r['lat_seek']}ms"
        )
    return "\n".join(lines)


EXPLANATION = [
    "# SMTC 播放器進度實測（M1 閘門）",
    "版本 v1.0 — 由 tools/measure_smtc.py 產生；每次執行 report 覆蓋本檔。",
    "# 方法",
    "以 50 毫秒間隔輪詢 Windows 媒體會話（SMTC）的 Timeline 與 PlaybackInfo，並同時記錄 SMTC 事件。連續播放期間，真實進度應以每秒 1 秒前進；每個播放區段（遇到暫停、拖動、切歌即切開）內，三種進度算法各自以「進度 = 經過秒數 + c」擬合，c 取該算法的中位數偏移。誤差為各樣本與這條線的差距。",
    "因 c 是擬合值，本測量量的是「回報的一致性」，無法量到回報進度與實際聲音之間的固定偏差（那需要擷取音訊，屬 L2）。",
    "# 指標",
    "進度更新：Position 或 LastUpdatedTime 改變的時刻；0.25 秒內的連續回報視為同一次（同一個狀態改變常觸發多筆）。更新間隔＝同一區段內相鄰兩次更新的時間差；播放器只在暫停、拖動時才回報的話，間隔欄會是空的，請看「更新次數／分」。",
    "直接取值誤差：把回報的 Position 當成「現在」的進度。播放器更新越稀疏，這個誤差越大（鋸齒狀）。",
    "外推誤差：Position + (現在 − LastUpdatedTime)，即 MediaWatcher 的算法。",
    "SyncEngine 誤差：只在收到新回報時重新對時、其餘時間本地推算的進度，與 1 倍速時鐘的差距。注意：播放器很少回報時，這條線幾乎不會被修正，數值會接近 0，不代表準確。",
    "對時校正量：每次收到新回報時，App 當下顯示所依據的進度與播放器新回報之間的差距，也就是重新對時前已經偏了多少。這是最接近非功能需求「歌詞時間誤差（相對播放器進度）≤ 100 毫秒」的指標；拖動造成的跳變（> 2 秒）不在區段內，不計入。",
    "漂移：外推進度對經過時間的斜率減 1，換算成每分鐘毫秒；正值代表回報跑得比時鐘快。",
    "回報延遲：腳本經 SMTC 下暫停／播放／拖動指令後，到輪詢看到狀態或進度改變的時間（僅在 --scripted 時有值）。",
    "少於 15 秒的區段不納入統計，但仍列在「區段明細」。",
    "# 判定",
    "M1 閘門（規格 v1.5）：Spotify 與 Chrome 的 YouTube Music 兩個播放器，對時校正量 p95 ≤ 100 毫秒。「M1 閘門判定」欄以公式計算；沒有對時資料時顯示「無資料」。KKBOX 不在範圍內。",
    "對時次數很少時（例如 Chrome 只在暫停、拖動時回報），判定只代表那幾次對時的結果，無法涵蓋播放中途緩衝卡住的情況。",
]


def write_xlsx(path: Path, rows: list[dict[str, object]], seg_rows: list[dict[str, object]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    font_name = "微軟正黑體"
    header_fill = PatternFill("solid", fgColor="DDE7F0")
    block_fill = PatternFill("solid", fgColor="FCE5CD")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def table(ws, columns, data):
        for c, (_, title, width, _) in enumerate(columns, 1):
            cell = ws.cell(row=1, column=c, value=title)
            cell.font = Font(name=font_name, bold=True, size=11)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
            ws.column_dimensions[get_column_letter(c)].width = width
        ws.row_dimensions[1].height = 36
        for r, row in enumerate(data, 2):
            longest = 1
            for c, (key, _, width, wrap) in enumerate(columns, 1):
                value = row.get(key)
                cell = ws.cell(row=r, column=c, value=value)
                cell.font = Font(name=font_name, size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=wrap, horizontal="left" if wrap else "center")
                cell.border = border
                if key == "status" and value != "已量測":
                    cell.fill = block_fill
                if wrap and value:
                    text = str(value)
                    units = sum(2 if ord(ch) > 0x2E80 else 1 for ch in text)
                    longest = max(longest, -(-units // max(1, int(width))))
            ws.row_dimensions[r].height = max(18, 15 * longest)
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(len(data) + 1, 2)}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:1"

    wb = Workbook()
    ws = wb.active
    ws.title = "摘要"
    table(ws, SUMMARY_COLUMNS, rows)

    # Formula statistics (Excel computes them on open).
    keys = [k for k, *_ in SUMMARY_COLUMNS]
    col = {k: get_column_letter(i + 1) for i, k in enumerate(keys)}
    last = len(rows) + 1
    for r, row in enumerate(rows, 2):
        if row.get("status") == "已量測":
            p95 = f"{col['corr_p95']}{r}"
            ws[f"{col['gate']}{r}"] = (
                f'=IF({p95}="","無資料",IF({p95}<={GATE_CORRECTION_P95_MS:g},"通過","未通過"))'
            )
    stats_row = last + 2
    stats = [
        ("已量測次數", f'=COUNTIFS({col["status"]}2:{col["status"]}{last},"已量測")'),
        ("BLOCK／未量測", f'=COUNTIFS({col["status"]}2:{col["status"]}{last},"<>已量測")'),
        ("M1 閘門通過次數", f'=COUNTIFS({col["gate"]}2:{col["gate"]}{last},"通過")'),
        ("M1 閘門未通過次數", f'=COUNTIFS({col["gate"]}2:{col["gate"]}{last},"未通過")'),
    ]
    for i, (label, formula) in enumerate(stats):
        label_cell = ws.cell(row=stats_row + i, column=1, value=label)
        label_cell.font = Font(name=font_name, size=10, bold=True)
        value_cell = ws.cell(row=stats_row + i, column=2, value=formula)
        value_cell.font = Font(name=font_name, size=10)
        value_cell.alignment = Alignment(horizontal="center")

    ws2 = wb.create_sheet("區段明細")
    table(ws2, SEGMENT_COLUMNS, seg_rows)

    ws3 = wb.create_sheet("說明")
    ws3.column_dimensions["A"].width = 110
    for r, para in enumerate(EXPLANATION, 1):
        cell = ws3.cell(row=r, column=1, value=para.lstrip("# ") if para.startswith("#") else para)
        cell.font = Font(name=font_name, size=11, bold=para.startswith("#"))
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        units = sum(2 if ord(ch) > 0x2E80 else 1 for ch in para)
        ws3.row_dimensions[r].height = max(18, 16 * -(-units // 105))

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
