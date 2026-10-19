"""Hit-rate regression over the test list (spec: acceptance and test strategy).

- candidates: songs from the local play history, for the user to pick into the list
- adopt: append the picked rows to the test list
- evaluate: query every listed song without the cache, then report the synced
  hit rate per language (English >= 85%, Chinese >= 75%) and scan everything
  that would be displayed for simplified characters (target: none)

Outputs carry song titles (listening history) and stay out of version
control; they never contain lyric text, only counts and scores.
"""

from __future__ import annotations

import asyncio
import csv
import statistics
import time
from collections.abc import Awaitable, Callable, Iterable, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from lyricfloat.diagnostics.problem_report import add_to_test_list
from lyricfloat.diagnostics.xlsx import Column, write_stats, write_table, write_text_sheet
from lyricfloat.lyrics.language import detect_from_metadata
from lyricfloat.models import LyricResult, ResultKind, TrackInfo

TARGETS = {"en": 0.85, "zh": 0.75}
LIST_TARGETS = {"en": 100, "zh": 80, "other": 20}
LANGUAGES = ("en", "zh", "other")
CATEGORIES = ("熱門", "冷門", "Live", "feat.", "重新母帶", "其他")
LANGUAGE_NAMES = {"en": "英文", "zh": "華語", "other": "其他"}


@dataclass(frozen=True)
class ListedSong:
    title: str
    artist: str
    album: str = ""
    duration: float = 0.0
    language: str = ""
    category: str = ""
    source_app: str = ""
    note: str = ""

    @property
    def key(self) -> tuple[str, str, str]:
        return (self.title.casefold(), self.artist.casefold(), self.album.casefold())

    def track(self) -> TrackInfo:
        return TrackInfo(self.title, self.artist, self.album, self.duration, self.source_app or "test-list")

    def lang(self) -> str:
        if self.language in LANGUAGES:
            return self.language
        return detect_from_metadata(self.title, self.artist).value


def _float(value: Any) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def load_test_list(path: Path) -> list[ListedSong]:
    try:
        with path.open(newline="", encoding="utf-8-sig") as fh:
            rows = list(csv.DictReader(fh))
    except FileNotFoundError:
        return []
    songs = []
    for row in rows:
        title = (row.get("title") or "").strip()
        if not title:
            continue
        songs.append(
            ListedSong(
                title=title,
                artist=(row.get("artist") or "").strip(),
                album=(row.get("album") or "").strip(),
                duration=_float(row.get("duration")),
                language=(row.get("language") or "").strip(),
                category=(row.get("category") or "").strip(),
                source_app=(row.get("source_app") or "").strip(),
                note=(row.get("note") or "").strip(),
            )
        )
    return songs


# -- candidates from play history -------------------------------------------------------


def build_candidates(history: Iterable[dict[str, Any]], listed: Iterable[ListedSong]) -> list[dict[str, Any]]:
    listed_keys = {s.key for s in listed}
    groups: dict[tuple[str, str, str], dict[str, Any]] = {}
    for item in history:
        key = (str(item.get("title", "")).casefold(), str(item.get("artist", "")).casefold(), str(item.get("album", "")).casefold())
        entry = groups.get(key)
        if entry is None:
            entry = groups[key] = {"plays": 0}
        entry["plays"] += 1
        entry.update(
            title=item.get("title", ""),
            artist=item.get("artist", ""),
            album=item.get("album", ""),
            duration=item.get("duration", 0.0),
            source_app=item.get("source_app", ""),
            last_played=item.get("at", ""),
            last_kind=item.get("kind", ""),
            last_source=item.get("source") or "",
            last_score=item.get("score"),
            language=item.get("language") or "",
            listed="是" if key in listed_keys else "否",
        )
    rows = list(groups.values())
    rows.sort(key=lambda r: (r["listed"] == "是", -r["plays"], r["title"]))
    return rows


CANDIDATE_COLUMNS = [
    Column("include", "納入", 8, fill_in=True, choices=("是", "否")),
    Column("language", "語言", 8, fill_in=True, choices=LANGUAGES),
    Column("category", "類型", 10, fill_in=True, choices=CATEGORIES),
    Column("title", "歌名", 28, wrap=True),
    Column("artist", "歌手", 22, wrap=True),
    Column("album", "專輯", 22, wrap=True),
    Column("duration", "時長 (s)", 9),
    Column("plays", "播放次數", 9),
    Column("last_played", "最近播放", 19),
    Column("last_kind", "上次結果", 12, colors={"not_found": "F4CCCC", "sources_failed": "F4CCCC", "plain_only": "FFF2CC"}),
    Column("last_source", "來源", 9),
    Column("last_score", "分數", 8),
    Column("source_app", "播放器", 12),
    Column("listed", "已在歌單", 9, colors={"是": "D9EAD3"}),
]

CANDIDATE_HELP = [
    "# 測試歌單候選",
    "由 App 的播放紀錄整理（每首歌只列一次，附播放次數）。規格目標：英文 100 首、華語 80 首、其他 20 首，涵蓋熱門、冷門、Live、feat. 與重新母帶版。",
    "# 怎麼用",
    "1. 在「候選」分頁的黃底欄位填寫：「納入」選「是」；「語言」確認或修正（en／zh／other）；「類型」選一項。",
    "2. 存檔後關閉 Excel，執行：python tools/hit_rate.py adopt",
    "3. 標成「是」的歌會加入測試歌單；已在歌單的歌不會重複加入。「摘要」分頁統計目前勾選數與目標的差距。",
    "上次結果為紅底（找不到、來源失敗）或黃底（只有純文字）的歌，特別適合加入，用來追蹤修正。",
]


def write_candidates_xlsx(path: Path, rows: Sequence[dict[str, Any]], listed: Sequence[ListedSong]) -> None:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "候選"
    prefilled = [{**r, "include": "否" if r["listed"] == "是" else "", "language": r.get("language", "")} for r in rows]
    write_table(ws, CANDIDATE_COLUMNS, prefilled, freeze="D2")

    col = {c.key: get_column_letter(i + 1) for i, c in enumerate(CANDIDATE_COLUMNS)}
    last = len(rows) + 1
    include, language = f"候選!{col['include']}2:{col['include']}{last}", f"候選!{col['language']}2:{col['language']}{last}"
    stats: list[tuple[str, Any]] = []
    for lang in LANGUAGES:
        name = LANGUAGE_NAMES[lang]
        first = len(stats) + 1  # row of "已在歌單" for this language
        stats += [
            (f"{name}：已在歌單", sum(1 for s in listed if s.lang() == lang)),
            (f"{name}：這次勾選納入", f'=COUNTIFS({include},"是",{language},"{lang}")'),
            (f"{name}：目標", LIST_TARGETS[lang]),
            (f"{name}：還差", f"=MAX(0,B{first + 2}-B{first}-B{first + 1})"),
        ]
    write_stats(wb.create_sheet("摘要"), 1, stats)
    write_text_sheet(wb.create_sheet("說明"), CANDIDATE_HELP)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def adopt_from_xlsx(xlsx: Path, test_list: Path | None = None) -> tuple[int, int]:
    """Append rows marked 納入=是 to the test list. Returns (added, already listed)."""
    from openpyxl import load_workbook

    ws = load_workbook(xlsx, read_only=True)["候選"]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows)]
    index = {title: header.index(title) for title in header}
    added = skipped = 0
    for values in rows:
        if values[index["納入"]] != "是":
            continue
        title = str(values[index["歌名"]] or "").strip()
        if not title:
            continue
        track = TrackInfo(
            title,
            str(values[index["歌手"]] or ""),
            str(values[index["專輯"]] or ""),
            _float(values[index["時長 (s)"]]),
            str(values[index["播放器"]] or ""),
        )
        language = str(values[index["語言"]] or "") or detect_from_metadata(track.title, track.artist).value
        category = str(values[index["類型"]] or "")
        if add_to_test_list(track, language, "", test_list, category):
            added += 1
        else:
            skipped += 1
    return added, skipped


# -- evaluation ---------------------------------------------------------------------


@dataclass
class Outcome:
    song: ListedSong
    kind: str = ""
    source: str = ""
    source_id: str = ""
    score: float | None = None
    l1_switched: bool = False
    lines: int = 0
    translated: int = 0
    simplified: int = 0
    first_ms: float | None = None
    error: str = ""

    @property
    def synced(self) -> bool:
        return self.kind == ResultKind.SYNCED.value


def count_simplified(result: LyricResult, shown: TrackInfo, is_simplified: Callable[[str], bool]) -> int:
    """Lines of displayed content still judged simplified: lyrics, translations, title and artist."""
    texts = [shown.title, shown.artist]
    if result.lyrics is not None:
        texts.extend(line.text for line in result.lyrics.lines)
        texts.extend(t for t in (result.lyrics.translations or ()) if t)
    return sum(1 for text in texts if is_simplified(text))


async def evaluate_song(song: ListedSong, resolver, pipeline) -> Outcome:
    outcome = Outcome(song)
    track = song.track()
    try:
        started = time.perf_counter()
        result = await resolver.resolve(track)
        outcome.first_ms = (time.perf_counter() - started) * 1000
        if resolver.needs_enrichment(result):
            result = await resolver.enrich(result)
    except Exception as exc:  # one broken song must not stop the run
        outcome.error = f"{type(exc).__name__}: {exc}"
        return outcome
    cand = result.candidate
    outcome.kind = result.kind.value
    outcome.source = cand.source if cand else ""
    outcome.source_id = cand.source_id if cand else ""
    outcome.score = result.score
    outcome.l1_switched = "l1_switched" in result.notes
    if result.lyrics is not None:
        outcome.lines = len(result.lyrics.lines)
        outcome.translated = sum(1 for t in (result.lyrics.translations or ()) if t)
    outcome.simplified = count_simplified(result, pipeline.display_track(track), pipeline.cc.is_simplified)
    return outcome


async def evaluate(
    songs: Sequence[ListedSong],
    resolver,
    pipeline,
    delay: float = 0.5,
    progress: Callable[[int, int, Outcome], Awaitable[None] | None] | None = None,
) -> list[Outcome]:
    outcomes = []
    for i, song in enumerate(songs, 1):
        outcome = await evaluate_song(song, resolver, pipeline)
        outcomes.append(outcome)
        if progress is not None:
            maybe = progress(i, len(songs), outcome)
            if asyncio.iscoroutine(maybe):
                await maybe
        if delay and i < len(songs):
            await asyncio.sleep(delay)  # be gentle with the sources
    return outcomes


def summarize(outcomes: Sequence[Outcome]) -> dict[str, Any]:
    by_lang: dict[str, dict[str, Any]] = {}
    for lang in LANGUAGES:
        group = [o for o in outcomes if o.song.lang() == lang]
        hits = sum(1 for o in group if o.synced)
        rate = hits / len(group) if group else None
        target = TARGETS.get(lang)
        by_lang[lang] = {
            "songs": len(group),
            "synced": hits,
            "rate": rate,
            "target": target,
            "passed": None if target is None or rate is None else rate >= target,
        }
    simplified = sum(o.simplified for o in outcomes)
    times = [o.first_ms for o in outcomes if o.first_ms is not None]
    return {
        "languages": by_lang,
        "simplified": simplified,
        "errors": sum(1 for o in outcomes if o.error),
        "first_ms_median": statistics.median(times) if times else None,
        "passed": simplified == 0 and all(v["passed"] is not False for v in by_lang.values()),
    }


REPORT_COLUMNS = [
    Column("title", "歌名", 26, wrap=True),
    Column("artist", "歌手", 20, wrap=True),
    Column("album", "專輯", 20, wrap=True),
    Column("duration", "時長 (s)", 9),
    Column("language", "語言", 8),
    Column("category", "類型", 9),
    Column("kind", "結果", 12, colors={"not_found": "F4CCCC", "sources_failed": "F4CCCC", "plain_only": "FFF2CC"}),
    Column("synced", "同步命中", 9, colors={"是": "D9EAD3", "否": "F4CCCC"}),
    Column("source", "來源", 9),
    Column("source_id", "編號", 12),
    Column("score", "分數", 8),
    Column("l1", "L1 換版本", 9),
    Column("lines", "歌詞行數", 9),
    Column("translated", "翻譯行數", 9),
    Column("simplified", "殘留簡體行數", 11, colors={str(n): "F4CCCC" for n in range(1, 500)}),
    Column("first_ms", "首行歌詞 (ms)", 11),
    Column("error", "錯誤", 24, wrap=True),
    Column("review", "人工備註", 24, wrap=True, fill_in=True),
]

REPORT_HELP = [
    "# 命中率回歸報告",
    "對測試歌單逐首查詢（不使用本機快取，反映來源本身的命中狀況），並跑完背景檢查（L1 時間軸共識、英文歌翻譯）。",
    "# 判準（規格：驗收標準）",
    "英文歌同步命中率 ≥ 85%、華語歌 ≥ 75%；「其他」只列數字不判定。同步命中＝取得同步歌詞（純文字、純音樂、找不到都不算）。",
    "顯示內容（歌詞、翻譯、歌名、歌手）經自動掃描，殘留簡體行數必須為 0。含假名或諺文的日韓文字不計入。",
    "首行歌詞 (ms)：從查詢到拿到第一份歌詞的時間，對照規格「切歌到顯示第一行 ≤ 1.5 秒（未命中快取）」，僅供參考。",
    "抽樣行時間軸誤差（中位數 ≤ 500 ms）由校對模式另外量測，不在本報告。",
    "# 注意",
    "本檔含歌名（屬聽歌紀錄），只存在本機、不進版控；不含任何歌詞文字。",
]


def write_report_xlsx(path: Path, outcomes: Sequence[Outcome]) -> None:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    rows = [
        {
            "title": o.song.title,
            "artist": o.song.artist,
            "album": o.song.album,
            "duration": o.song.duration,
            "language": o.song.lang(),
            "category": o.song.category,
            "kind": o.kind or "error",
            "synced": "是" if o.synced else "否",
            "source": o.source,
            "source_id": o.source_id,
            "score": o.score,
            "l1": "是" if o.l1_switched else "",
            "lines": o.lines,
            "translated": o.translated,
            "simplified": o.simplified,
            "first_ms": None if o.first_ms is None else round(o.first_ms),
            "error": o.error,
        }
        for o in outcomes
    ]
    wb = Workbook()
    summary = wb.active
    summary.title = "摘要"
    detail = wb.create_sheet("明細")
    write_table(detail, REPORT_COLUMNS, rows, freeze="C2")

    col = {c.key: get_column_letter(i + 1) for i, c in enumerate(REPORT_COLUMNS)}
    last = len(rows) + 1
    rng = lambda key: f"明細!{col[key]}2:{col[key]}{last}"  # noqa: E731
    stats: list[tuple[str, Any]] = []
    for lang in LANGUAGES:
        name = LANGUAGE_NAMES[lang]
        songs = len(stats) + 1  # row numbers in column B of this sheet
        stats += [
            (f"{name}：歌曲數", f'=COUNTIFS({rng("language")},"{lang}")'),
            (f"{name}：同步命中", f'=COUNTIFS({rng("language")},"{lang}",{rng("synced")},"是")'),
            (f"{name}：命中率", f'=IF(B{songs}=0,"—",B{songs + 1}/B{songs})'),
        ]
        if lang in TARGETS:
            stats += [
                (f"{name}：目標", TARGETS[lang]),
                (f"{name}：判定", f'=IF(B{songs}=0,"無資料",IF(B{songs + 2}>=B{songs + 3},"通過","未通過"))'),
            ]
    simplified = len(stats) + 1
    stats += [
        ("殘留簡體行數（全部）", f"=SUM({rng('simplified')})"),
        ("簡體判定", f'=IF(B{simplified}=0,"通過","未通過")'),
        ("查詢出錯的歌", f'=COUNTIFS({rng("kind")},"error")'),
        ("首行歌詞時間中位數 (ms)", f"=MEDIAN({rng('first_ms')})"),
    ]
    write_stats(summary, 1, stats)
    for i, (label, _value) in enumerate(stats, 1):
        if label.endswith(("命中率", "目標")):
            summary.cell(row=i, column=2).number_format = "0.0%"
    write_text_sheet(wb.create_sheet("說明"), REPORT_HELP)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
