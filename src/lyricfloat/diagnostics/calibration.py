"""Calibration mode (SPEC v1.6): measure timeline error by hand.

While listening, the user presses a hotkey the moment a line is actually
sung. Each mark records press position minus the displayed start of the
nearest line (within 3 s): positive means the lyrics showed up early,
negative late. Five marks per song, about twenty songs; acceptance is a
median absolute error of at most 500 ms. Human reaction time (roughly
150-250 ms) is included and not subtracted, so the figure is conservative.
"""

from __future__ import annotations

import datetime as dt
import json
import statistics
from bisect import bisect_left
from collections.abc import Sequence
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from lyricfloat.logging_setup import log_dir

MARK_WINDOW = 3.0  # seconds either side of the press
MARKS_PER_SONG = 5
TARGET_MEDIAN_MS = 500.0


def calibration_path() -> Path:
    return log_dir() / "calibration.jsonl"


def nearest_line(starts: Sequence[float], position: float, window: float = MARK_WINDOW) -> int | None:
    """Index of the line start closest to `position`, if one lies within `window` seconds."""
    if not starts:
        return None
    i = bisect_left(starts, position)
    best = None
    for j in (i - 1, i):
        if 0 <= j < len(starts) and abs(starts[j] - position) <= window:
            if best is None or abs(starts[j] - position) < abs(starts[best] - position):
                best = j
    return best


@dataclass(frozen=True)
class Mark:
    at: str
    title: str
    artist: str
    source: str
    source_id: str
    line_index: int
    line_start: float  # displayed start, offset included
    offset: float
    position: float  # where playback was when the key was pressed
    delta_ms: float  # position - line_start


class CalibrationRecorder:
    def __init__(self, path: Path | None = None) -> None:
        self.path = path or calibration_path()
        self._counts: dict[tuple[str, str], int] = {}
        for mark in read_marks(self.path):
            key = (mark["title"], mark["artist"])
            self._counts[key] = self._counts.get(key, 0) + 1

    def count(self, title: str, artist: str) -> int:
        return self._counts.get((title, artist), 0)

    def record(
        self,
        title: str,
        artist: str,
        source: str,
        source_id: str,
        starts: Sequence[float],
        offset: float,
        position: float,
    ) -> Mark | None:
        """Store a mark; None when no line is near or the song already has five."""
        if self.count(title, artist) >= MARKS_PER_SONG:
            return None
        index = nearest_line(starts, position)
        if index is None:
            return None
        mark = Mark(
            at=dt.datetime.now().isoformat(timespec="seconds"),
            title=title,
            artist=artist,
            source=source,
            source_id=source_id,
            line_index=index,
            line_start=round(starts[index], 3),
            offset=offset,
            position=round(position, 3),
            delta_ms=round((position - starts[index]) * 1000, 1),
        )
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with self.path.open("a", encoding="utf-8") as fh:
            fh.write(json.dumps(asdict(mark), ensure_ascii=False) + "\n")
        self._counts[(title, artist)] = self.count(title, artist) + 1
        return mark


def read_marks(path: Path | None = None) -> list[dict[str, Any]]:
    path = path or calibration_path()
    try:
        lines = path.read_text(encoding="utf-8").splitlines()
    except FileNotFoundError:
        return []
    marks = []
    for line in lines:
        try:
            item = json.loads(line)
        except ValueError:
            continue
        if isinstance(item, dict) and "delta_ms" in item:
            marks.append(item)
    return marks


def summarize_marks(marks: Sequence[dict[str, Any]]) -> dict[str, Any]:
    errors = [abs(float(m["delta_ms"])) for m in marks]
    songs = {(m["title"], m["artist"]) for m in marks}
    median = statistics.median(errors) if errors else None
    signed = statistics.median(float(m["delta_ms"]) for m in marks) if marks else None
    return {
        "marks": len(marks),
        "songs": len(songs),
        "median_abs_ms": median,
        "median_signed_ms": signed,
        "passed": None if median is None else median <= TARGET_MEDIAN_MS,
    }


MARK_HELP = [
    "# 校對結果（抽樣行時間軸誤差）",
    "校對模式下，聽到每句實際開唱的瞬間按快捷鍵；誤差＝按鍵當下的播放位置 − 該行的顯示時間（含使用者偏移）。正值：歌詞比歌聲早出現；負值：歌詞比歌聲晚。",
    "規格目標：抽 20 首、每首 5 行，誤差絕對值的中位數 ≤ 500 毫秒。",
    "# 注意",
    "按鍵含人的反應時間（約 150–250 毫秒），本表沒有扣除，所以結果偏保守；帶正負的中位數若明顯為正，多半就是反應時間。",
    "本檔含歌名（屬聽歌紀錄），只存在本機、不進版控；不含任何歌詞文字。",
]


def write_marks_xlsx(path: Path, marks: Sequence[dict[str, Any]]) -> None:
    from openpyxl import Workbook

    from lyricfloat.diagnostics.xlsx import Column, write_stats, write_table, write_text_sheet

    columns = [
        Column("at", "時間", 19),
        Column("title", "歌名", 26, wrap=True),
        Column("artist", "歌手", 20, wrap=True),
        Column("source", "來源", 9),
        Column("source_id", "編號", 12),
        Column("line_index", "行號", 7),
        Column("line_start", "顯示時間 (s)", 11),
        Column("offset", "偏移 (s)", 9),
        Column("position", "按鍵位置 (s)", 11),
        Column("delta_ms", "誤差 (ms)", 10),
        Column("abs_ms", "誤差絕對值 (ms)", 12),
        Column("review", "人工備註", 24, wrap=True, fill_in=True),
    ]
    rows = [{**m, "abs_ms": f"=ABS(J{i})"} for i, m in enumerate(marks, 2)]
    wb = Workbook()
    summary = wb.active
    summary.title = "摘要"
    write_table(wb.create_sheet("標記"), columns, rows, freeze="C2")
    last = len(rows) + 1
    write_stats(
        summary,
        1,
        [
            ("標記數", f"=COUNT(標記!J2:J{last})"),
            ("歌曲數", len({(m["title"], m["artist"]) for m in marks})),
            ("誤差絕對值中位數 (ms)", f'=IF(B1=0,"—",MEDIAN(標記!K2:K{last}))'),
            ("帶正負的誤差中位數 (ms)", f'=IF(B1=0,"—",MEDIAN(標記!J2:J{last}))'),
            ("目標 (ms)", TARGET_MEDIAN_MS),
            ("判定", '=IF(B1=0,"無資料",IF(B3<=B5,"通過","未通過"))'),
        ],
    )
    write_text_sheet(wb.create_sheet("說明"), MARK_HELP)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
