"""Browser title self-test (SPEC v1.50): songs whose lyrics are known, dressed up as the
titles YouTube uploads give them, resolved as Chrome would report them.

Each song is first resolved from its clean metadata, as a music app reports it: that
result is the answer. Every variant of the song must then lead to the same song. The
templates copy title conventions seen on YouTube (official channels, Topic channels,
lyric channels, shows, K-pop stages, Japanese brackets, simplified/traditional uploads).

Outputs carry song titles (listening history) and stay out of version control; they
never contain lyric text. The lyric snippets in templates are made up.
"""

from __future__ import annotations

import asyncio
import re
from collections.abc import Awaitable, Callable, Iterable, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from lyricfloat.diagnostics.xlsx import Column, write_stats, write_table, write_text_sheet
from lyricfloat.lyrics.identify import song_key
from lyricfloat.lyrics.normalize import other_language, split_artists
from lyricfloat.lyrics.sources.base import LyricSource
from lyricfloat.models import Candidate, LyricResult, ResultKind, TrackInfo

BROWSER_APP = "Chrome"
CLEAN_APP = "Spotify.exe"
GROUPS = ("zh", "jp", "kr", "en")
GROUP_NAMES = {"zh": "華語", "jp": "日文", "kr": "韓文", "en": "英文"}

OK = "正確"
OK_PLAIN = "正確（無同步）"
NOT_FOUND = "找不到"
WRONG = "抓錯歌"
SKIPPED = "不是歌"
FAILED = "來源失敗"
VERDICTS = (OK, OK_PLAIN, NOT_FOUND, WRONG, SKIPPED, FAILED)

_FEAT = re.compile(r"\s*[(\[](?:feat\.?|ft\.?|with)\s*([^()\[\]]+)[)\]]", re.IGNORECASE)
_KANA = re.compile(r"[ぁ-ゟ゠-ヿ]")
_HAN = re.compile(r"[㐀-䶿一-鿿豈-﫿]")
_SPACES = re.compile(r"\s+")


@dataclass(frozen=True)
class Song:
    title: str
    artist: str
    duration: float
    language: str  # the app's metadata language: en / zh / other
    album: str = ""

    def track(self) -> TrackInfo:
        # The album keeps Spotify's ad rule (no album, under a minute) away from short songs.
        return TrackInfo(self.title, self.artist, self.album or self.title, self.duration, CLEAN_APP)


@dataclass(frozen=True)
class Parts:
    """A song as video titles name it: the bare title (no " - Live", no feat.), the artist."""

    title: str
    artist: str
    feat: str
    duration: float
    group: str


@dataclass(frozen=True)
class Variant:
    template: str
    group: str
    title: str
    channel: str
    duration: float

    def track(self) -> TrackInfo:
        return TrackInfo(self.title, self.channel, "", self.duration, BROWSER_APP)


Convert = Callable[[str, str], str]  # (opencc config, text) -> text


@dataclass(frozen=True)
class Template:
    name: str
    label: str
    groups: tuple[str, ...]
    make: Callable[[Parts, Convert], tuple[str, str, float] | None]


def group_of(song: Song) -> str:
    """Which title conventions fit: Chinese, Japanese (kana, or Han in a song the app does
    not call Chinese), Korean (other songs the app does not call English or Chinese), English."""
    text = f"{song.title} {song.artist}"
    if song.language == "zh":
        return "zh"
    if song.language == "en":
        return "en"
    if _KANA.search(text) or _HAN.search(text):
        return "jp"
    return "kr"


def parts_of(song: Song) -> Parts:
    title = song.title.split(" - ")[0].strip()  # Spotify's " - Live", " - 電視劇《…》片尾曲"
    feat = ""
    m = _FEAT.search(title)
    if m:
        feat = m.group(1).strip()
        title = (title[: m.start()] + title[m.end():]).strip()
    return Parts(title, song.artist.strip(), feat, song.duration, group_of(song))


def _squash(text: str) -> str:
    return _SPACES.sub("", text)


def _t(fmt: str, channel: str, extra: float = 0.0):
    """A template from format strings: {t} title, {a} artist, {f} featured artist."""

    def make(p: Parts, _cc: Convert):
        if "{f}" in fmt and not p.feat:
            return None
        fields = {"t": p.title, "a": p.artist, "f": p.feat, "A": _squash(p.artist), "T": _squash(p.title),
                  "al": p.artist.casefold(), "tl": p.title.casefold()}
        return fmt.format(**fields), channel.format(**fields), p.duration + extra

    return make


def _converted(config: str, fmt: str, channel: str):
    def make(p: Parts, cc: Convert):
        title, chan, duration = _t(fmt, channel)(p, cc)
        return cc(config, title), chan, duration

    return make


ALL = GROUPS
TEMPLATES: tuple[Template, ...] = (
    # Any language
    Template("mv", "歌手 - 歌名 (Official Music Video)，官方頻道，MV 前後多 12 秒", ALL,
             _t("{a} - {t} (Official Music Video)", "{a}", 12)),
    Template("mv_bracket", "歌手 - 歌名 [Official MV]", ALL, _t("{a} - {t} [Official MV]", "{a} Official", 8)),
    Template("topic", "Topic 頻道：只有歌名，頻道「歌手 - Topic」", ALL, _t("{t}", "{a} - Topic")),
    Template("title_only", "官方頻道只寫歌名 (Official Music Video)", ALL, _t("{t} (Official Music Video)", "{a}", 10)),
    Template("audio", "歌手 - 歌名 (Official Audio)", ALL, _t("{a} - {t} (Official Audio)", "{a}")),
    Template("lyrics", "歌詞頻道：歌手 - 歌名 (Lyrics)", ALL, _t("{a} - {t} (Lyrics)", "7clouds")),
    Template("lyrics_reversed", "歌詞頻道：歌名 - 歌手 (Lyrics)", ALL, _t("{t} - {a} (Lyrics)", "Dan Music")),
    Template("lyrics_pipe", "歌詞頻道：歌名 | 歌手 (Lyric Video)", ALL, _t("{t} | {a} (Lyric Video)", "Lyric Vibes")),
    Template("visualizer", "歌手 – 歌名 (Visualizer)（長破折號）", ALL, _t("{a} – {t} (Visualizer)", "{a}")),
    Template("slash", "歌名 / 歌手", ALL, _t("{t} / {a}", "{a}")),
    Template("live", "歌手 - 歌名 (Live)，多 6 秒", ALL, _t("{a} - {t} (Live)", "{a}", 6)),
    Template("mv_edit", "唱片公司剪短的 MV：歌手 [歌名] Official MV，少 8 秒", ALL,
             _t("{a} [{t}] Official MV", "華研國際", -8)),
    Template("feat", "歌手 - 歌名 (feat. 客串) [Official Music Video]", ALL,
             _t("{a} - {t} (feat. {f}) [Official Music Video]", "{a}", 10)),
    # Chinese
    Template("zh_book", "歌手《歌名》Official Music Video，唱片公司頻道", ("zh",),
             _t("{a}《{t}》Official Music Video", "華納音樂 Warner Music Taiwan", 10)),
    Template("zh_lenticular", "歌手【歌名】Official MV，唱片公司頻道", ("zh",),
             _t("{a}【{t}】Official MV", "滾石唱片 ROCK RECORDS", 10)),
    Template("zh_hq", "歌手 - 歌名【高音質|動態歌詞 Lyrics】", ("zh",),
             _t("{a} - {t}【高音質|動態歌詞 Lyrics】", "動態歌詞 Lyrics Channel")),
    Template("zh_snippet", "歌手 - 歌名『一句歌詞』【動態歌詞Lyrics】", ("zh",),
             _t("{a} - {t}『我們都在等一場雨停』【動態歌詞Lyrics】", "聽歌狂人")),
    Template("zh_snippet_reversed", "歌名 - 歌手「一句歌詞」【動態歌詞/Pinyin Lyrics】", ("zh",),
             _t("{t} - {a}「我們都在等一場雨停」【動態歌詞/Pinyin Lyrics】", "動態歌詞-Pinyin Lyrics")),
    Template("zh_simplified", "簡體上傳：歌手 - 歌名【动态歌词】", ("zh",),
             _converted("t2s", "{a} - {t}【动态歌词】", "抖音热歌榜")),
    Template("zh_traditional", "繁體上傳：歌名 - 歌手【動態歌詞】", ("zh",),
             _converted("s2t", "{t} - {a}【動態歌詞】", "華語音樂台")),
    Template("zh_show", "節目純享：【纯享】歌手《歌名》 | 節目 EP03 | iQIYI", ("zh",),
             _t("【纯享】{a}《{t}》 | 说唱巅峰对决 EP03 | iQIYI", "iQIYI 爱奇艺", 4)),
    Template("zh_lossless", "歌手《歌名》【动态歌词Lyrics】【无损音质】", ("zh",),
             _t("{a}《{t}》【动态歌词Lyrics】【无损音质】", "音乐热播")),
    Template("zh_hashtags", "歌手 - 歌名 #歌手 #歌名 #中文歌", ("zh",), _t("{a} - {t} #{A} #{T} #中文歌", "{a}")),
    Template("zh_pinyin", "歌名 - 歌手 (拼音歌詞 Pinyin Lyrics)", ("zh",),
             _t("{t} - {a} (拼音歌詞 Pinyin Lyrics)", "Pinyin Lyrics")),
    # Korean
    Template("kr_mv", "歌手 '歌名' M/V", ("kr",), _t("{a} '{t}' M/V", "{a}", 10)),
    Template("kr_label", "歌手 '歌名' Official MV，經紀公司頻道", ("kr",), _t("{a} '{t}' Official MV", "HYBE LABELS", 10)),
    Template("kr_1thek", "[MV] 歌手 _ 歌名（1theK）", ("kr",), _t("[MV] {a} _ {t}", "1theK (원더케이)", 5)),
    Template("kr_inkigayo", "人氣歌謠舞台：歌手 - 歌名 @인기가요 inkigayo 20241006", ("kr",),
             _t("{a} - {t} @인기가요 inkigayo 20241006", "SBS KPOP", 3)),
    Template("kr_mcountdown", "M Countdown 舞台：歌手 - 歌名 #엠카운트다운 EP.880 | Mnet 241010 방송", ("kr",),
             _t("{a} - {t} #엠카운트다운 EP.880 | Mnet 241010 방송", "Mnet K-POP", 3)),
    Template("kr_color_coded", "歌詞頻道：歌手 - 歌名 (Color Coded Lyrics Eng/Rom/Han/가사)", ("kr",),
             _t("{a} - {t} (Color Coded Lyrics Eng/Rom/Han/가사)", "Sing Lyrics")),
    # Japanese
    Template("jp_mv", "歌手「歌名」Official Music Video", ("jp",), _t("{a}「{t}」Official Music Video", "{a}", 10)),
    Template("jp_mv_slash", "【MV】歌名 / 歌手", ("jp",), _t("【MV】{t} / {a}", "{a}", 5)),
    Template("jp_first_take", "歌手 - 歌名 / THE FIRST TAKE，開頭多 25 秒", ("jp",),
             _t("{a} - {t} / THE FIRST TAKE", "THE FIRST TAKE", 25)),
    Template("jp_lyrics", "歌手「歌名」歌詞付き Lyrics", ("jp",), _t("{a}「{t}」歌詞付き Lyrics", "Lyrics JP")),
    # English
    Template("en_vevo", "歌手 - 歌名 (Official Video)，VEVO 頻道", ("en",), _t("{a} - {t} (Official Video)", "{A}VEVO", 15)),
    Template("en_snippet", "歌手 - 歌名 (Lyrics) \"一句歌詞\"", ("en",),
             _t('{a} - {t} (Lyrics) "and I keep running back to you"', "Vibe Music")),
    Template("en_lowercase", "全小寫：歌手 - 歌名 (lyrics)", ("en",), _t("{al} - {tl} (lyrics)", "lyric vibes")),
)


def variants_of(song: Song, cc: Convert, templates: Sequence[Template] = TEMPLATES) -> list[Variant]:
    p = parts_of(song)
    out = []
    for tpl in templates:
        if p.group not in tpl.groups:
            continue
        made = tpl.make(p, cc)
        if made is not None:
            title, channel, duration = made
            out.append(Variant(tpl.name, p.group, title, channel, duration))
    return out


# -- judging --------------------------------------------------------------------------


Fold = Callable[[str], str]


def same_song(expected: Candidate, got: Candidate, fold: Fold = lambda s: s) -> bool:
    """The same entry, or the same song (title however spelt, one artist in common)."""
    if (expected.source, expected.source_id) == (got.source, got.source_id):
        return True
    if other_language(expected.title, got.title):  # "Love Talk (English Version)" is other lyrics (v1.70)
        return False
    ours, theirs = song_key(expected.title, fold), song_key(got.title, fold)
    # A version's title may say more: "CRY FOR ME (live) FULL Ver." is TWICE's CRY FOR ME.
    if ours != theirs and not (len(ours) >= 4 and theirs.startswith(ours)):
        return False
    # 163 braces is 163braces; "NICKTHEREAL 周湯豪" is 周汤豪 with his English name
    names = [_squash(fold(n)) for n in split_artists(expected.artist)]
    others = [_squash(fold(x)) for x in split_artists(got.artist)]
    if any(len(m) >= 2 and len(n) >= 2 and (m in n or n in m) for n in names for m in others):
        return True
    # One stage name in both: NINEONE赵馨玥 and NINEONE#乃万
    return bool(_latin_words(expected.artist) & _latin_words(got.artist))


_LATIN_NAME = re.compile(r"[a-z]{4,}")
_COMMON_WORDS = {"band", "music", "official", "group", "crew", "project", "orchestra", "choir", "records",
                 "sound", "boys", "girls", "cover", "covers", "remix", "lyrics", "topic", "vevo", "feat", "with"}


def _latin_words(artist: str) -> set[str]:
    return set(_LATIN_NAME.findall(artist.casefold())) - _COMMON_WORDS


_NOT_WORD = re.compile(r"[\W_]+")
MIN_LINE = 4  # characters: "oh", "yeah" lines are in every song
SAME_LYRICS = 0.5  # share of the shorter side's distinct lines found in the other


def _line_set(result: LyricResult) -> set[str]:
    if result.lyrics is None:
        return set()
    lines = (_NOT_WORD.sub("", line.text).casefold() for line in result.lyrics.lines)
    return {line for line in lines if len(line) >= MIN_LINE}


def same_lyrics(expected: LyricResult, got: LyricResult) -> bool:
    """The lines on screen are the song's: LRCLIB also files songs under a video's title
    ("Taylor Swift - Style (Taylor's Version) (Lyric Video)") or with title and artist swapped."""
    a, b = _line_set(expected), _line_set(got)
    return bool(a and b) and len(a & b) >= SAME_LYRICS * min(len(a), len(b))


BY_ENTRY = "歌名歌手相同"
BY_LYRICS = "歌詞相同（別的條目）"


def judge(expected: LyricResult, got: LyricResult, fold: Fold = lambda s: s) -> tuple[str, str]:
    """(verdict, how the song was recognised as the answer)."""
    if got.kind in (ResultKind.NOT_A_SONG, ResultKind.SKIPPED_AD):
        return SKIPPED, ""
    if got.kind is ResultKind.SOURCES_FAILED:
        return FAILED, ""
    if got.candidate is None or expected.candidate is None:
        return (FAILED if "source_failed" in got.notes else NOT_FOUND), ""
    if same_song(expected.candidate, got.candidate, fold):
        how = BY_ENTRY
    elif same_lyrics(expected, got):
        how = BY_LYRICS
    else:
        return WRONG, ""
    if expected.kind is ResultKind.SYNCED and got.kind is not ResultKind.SYNCED:
        return OK_PLAIN, how
    return OK, how


# -- running ---------------------------------------------------------------------------


class MemoSource(LyricSource):
    """A source that asks the network once per distinct request: many variants of a song
    send the same query. Failures are not remembered."""

    def __init__(self, inner: LyricSource) -> None:
        self.inner = inner
        self.name = inner.name  # type: ignore[misc]
        self.simplified = inner.simplified  # type: ignore[misc]
        self.lyrics_inline = inner.lyrics_inline  # type: ignore[misc]
        self._memo: dict[tuple, Any] = {}
        self.requests = 0

    async def _once(self, key: tuple, call: Callable[[], Awaitable[Any]]) -> Any:
        if key not in self._memo:
            self.requests += 1
            self._memo[key] = await call()
        return self._memo[key]

    async def lookup(self, track: TrackInfo) -> Candidate | None:
        return await self._once(("lookup", track), lambda: self.inner.lookup(track))

    async def search(self, track: TrackInfo) -> list[Candidate]:
        return await self._once(("search", track), lambda: self.inner.search(track))

    async def search_text(self, text: str) -> list[Candidate]:
        return await self._once(("text", text), lambda: self.inner.search_text(text))

    async def fetch_lyrics(self, candidate: Candidate) -> Candidate:
        return await self._once(("fetch", candidate.source_id), lambda: self.inner.fetch_lyrics(candidate))

    async def aclose(self) -> None:
        await self.inner.aclose()


@dataclass(frozen=True)
class Outcome:
    song: Song
    variant: Variant
    verdict: str
    how: str
    expected: str
    got: str
    kind: str


def _describe(result: LyricResult) -> str:
    c = result.candidate
    return f"{c.title} — {c.artist}（{c.source} {c.source_id}，{c.duration:.0f} 秒）" if c else ""


async def evaluate(
    songs: Iterable[Song],
    resolver,  # LyricResolver with a NullCache: every variant asks the sources
    cc: Convert,
    fold: Fold,
    delay: float = 0.3,
    progress: Callable[[int, int, Outcome | None, Song], None] | None = None,
    templates: Sequence[Template] = TEMPLATES,
) -> tuple[list[Outcome], list[tuple[Song, str]]]:
    """Variants of every song with an answer; songs without one come back with the reason."""
    outcomes: list[Outcome] = []
    skipped: list[tuple[Song, str]] = []
    plan = [(song, variants_of(song, cc, templates)) for song in songs]
    total = sum(len(v) for _, v in plan)
    done = 0
    for song, variants in plan:
        expected = await resolver.resolve(song.track())
        if expected.kind not in (ResultKind.SYNCED, ResultKind.PLAIN_ONLY):
            skipped.append((song, f"乾淨歌名也沒有歌詞（{expected.kind.value}）"))
            done += len(variants)
            if progress:
                progress(done, total, None, song)
            continue
        for v in variants:
            resolver.forget_misses()  # each variant on its own: no "not found this session"
            got = await resolver.resolve(v.track())
            o = Outcome(song, v, *judge(expected, got, fold), _describe(expected), _describe(got), got.kind.value)
            outcomes.append(o)
            done += 1
            if progress:
                progress(done, total, o, song)
            await asyncio.sleep(delay)
    return outcomes, skipped


def pick_songs(history: Iterable[dict[str, Any]], per_group: int = 0, groups: Sequence[str] = GROUPS) -> list[Song]:
    """Distinct songs the app found lyrics for from a music app (not a browser), newest
    play first, at most `per_group` per title convention (0: all)."""
    seen: dict[tuple[str, str], Song] = {}
    for row in reversed(list(history)):
        app = str(row.get("source_app", "")).casefold()
        if any(b in app for b in ("chrome", "msedge", "firefox", "opera", "brave", "demo")):
            continue
        if row.get("kind") not in ("synced", "plain_only") or not row.get("title") or not row.get("artist"):
            continue
        key = (row["title"].casefold(), row["artist"].casefold())
        if key not in seen and row["artist"] != "LyricFloat":
            seen[key] = Song(row["title"], row["artist"], float(row.get("duration") or 0), row.get("language", ""),
                             row.get("album", ""))
    counts: dict[str, int] = {}
    out = []
    for song in seen.values():
        g = group_of(song)
        if g not in groups or (per_group and counts.get(g, 0) >= per_group):
            continue
        counts[g] = counts.get(g, 0) + 1
        out.append(song)
    return out


# -- report ----------------------------------------------------------------------------

_VERDICT_COLORS = {OK: "E2EFDA", OK_PLAIN: "FFF2CC", NOT_FOUND: "FCE4D6", WRONG: "F8CBAD", FAILED: "D9D9D9"}

DETAIL_COLUMNS = [
    Column("verdict", "結果", 12, colors=_VERDICT_COLORS),
    Column("group", "語言", 8),
    Column("template", "標題格式", 16),
    Column("title", "模擬影片標題", 48, wrap=True),
    Column("channel", "頻道（歌手欄）", 22, wrap=True),
    Column("duration", "影片秒數", 9),
    Column("song", "歌（乾淨歌名 — 歌手）", 28, wrap=True),
    Column("expected", "答案", 34, wrap=True),
    Column("got", "影片解析結果", 34, wrap=True),
    Column("how", "比對方式", 14, wrap=True),
    Column("note", "備註", 24, wrap=True, fill_in=True),
]

HELP = [
    "# 瀏覽器影片標題自測（SPEC v1.50）",
    "每首歌先用乾淨的歌名與歌手（像 Spotify 回報的）解析，得到的歌當答案；再把它套進 YouTube 常見的標題格式，"
    "以 Chrome 影片的身分（標題＝影片標題、歌手欄＝頻道名）重新解析，不經快取。",
    "結果：正確＝找到同一首歌（同一筆、歌名相同且有共同歌手，或歌詞有一半以上的行相同：LRCLIB 也有用影片標題或歌名歌手對調上傳的條目）；正確（無同步）＝同一首但只有純文字；"
    "找不到；抓錯歌＝顯示了別首歌的歌詞（最嚴重）；不是歌＝被判定不查詢；來源失敗＝網路或來源沒有回應，不算程式問題。",
    "模擬標題裡的「一句歌詞」是編的，不是真的歌詞；檔案只有歌名與歌手，沒有歌詞內容。",
    "# 各分頁",
    "摘要：各結果的數量（公式）。各格式：每種標題格式的正確率，先看正確率低的格式。明細：每個模擬標題一列，可篩選。"
    "沒有答案的歌：乾淨歌名也找不到歌詞的歌，不產生變體。",
]


def template_rows(outcomes: Sequence[Outcome], templates: Sequence[Template] = TEMPLATES) -> list[dict[str, Any]]:
    rows = []
    used = {o.variant.template for o in outcomes}
    for tpl in templates:
        if tpl.name in used:
            rows.append({"template": tpl.name, "label": tpl.label, "groups": "、".join(GROUP_NAMES[g] for g in tpl.groups)})
    return rows


def write_report_xlsx(path: Path, outcomes: Sequence[Outcome], skipped: Sequence[tuple[Song, str]]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "摘要"
    n = max(len(outcomes) + 1, 2)
    stats: list[tuple[str, Any]] = [("模擬標題總數", f"=COUNTA(明細!A2:A{n})")]
    stats += [(v, f'=COUNTIFS(明細!A2:A{n},"{v}")') for v in VERDICTS]
    stats.append(("正確率（正確＋無同步）", "=IFERROR((B2+B3)/B1,0)"))
    stats.append(("測試歌曲數", len({(o.song.title, o.song.artist) for o in outcomes})))
    stats.append(("沒有答案的歌", len(skipped)))
    write_stats(ws, 1, stats)
    ws.cell(row=8, column=2).number_format = "0.0%"

    fmt = wb.create_sheet("各格式")
    rows = template_rows(outcomes)
    write_table(fmt, [Column("template", "標題格式", 18), Column("label", "說明", 56, wrap=True), Column("groups", "適用", 18),
                      Column("total", "總數", 8), Column("ok", "正確", 8), Column("not_found", "找不到", 8),
                      Column("wrong", "抓錯歌", 8), Column("rate", "正確率", 10)], rows)
    for r in range(2, len(rows) + 2):
        a = f"$A{r}"
        fmt[f"D{r}"] = f"=COUNTIFS(明細!$C$2:$C${n},{a})"
        fmt[f"E{r}"] = f'=COUNTIFS(明細!$C$2:$C${n},{a},明細!$A$2:$A${n},"{OK}")+COUNTIFS(明細!$C$2:$C${n},{a},明細!$A$2:$A${n},"{OK_PLAIN}")'
        fmt[f"F{r}"] = f'=COUNTIFS(明細!$C$2:$C${n},{a},明細!$A$2:$A${n},"{NOT_FOUND}")'
        fmt[f"G{r}"] = f'=COUNTIFS(明細!$C$2:$C${n},{a},明細!$A$2:$A${n},"{WRONG}")'
        fmt[f"H{r}"] = f"=IFERROR(E{r}/D{r},0)"
        fmt[f"H{r}"].number_format = "0.0%"

    detail = wb.create_sheet("明細")
    write_table(detail, DETAIL_COLUMNS, [
        {"verdict": o.verdict, "group": GROUP_NAMES[o.variant.group], "template": o.variant.template, "title": o.variant.title,
         "channel": o.variant.channel, "duration": round(o.variant.duration, 1), "song": f"{o.song.title} — {o.song.artist}",
         "expected": o.expected, "got": o.got or o.kind, "how": o.how}
        for o in outcomes
    ])

    miss = wb.create_sheet("沒有答案的歌")
    write_table(miss, [Column("song", "歌", 40, wrap=True), Column("reason", "原因", 40, wrap=True)],
                [{"song": f"{s.title} — {s.artist}", "reason": why} for s, why in skipped])

    write_text_sheet(wb.create_sheet("說明"), HELP)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def outcome_json(o: Outcome) -> dict[str, Any]:
    v = o.variant
    return {"verdict": o.verdict, "how": o.how, "group": v.group, "template": v.template, "title": v.title, "channel": v.channel,
            "duration": v.duration, "song": o.song.title, "artist": o.song.artist, "expected": o.expected, "got": o.got,
            "kind": o.kind}


__all__ = [
    "MemoSource", "Outcome", "Song", "TEMPLATES", "Template", "Variant", "VERDICTS", "evaluate", "group_of", "judge",
    "outcome_json", "parts_of", "pick_songs", "same_song", "variants_of", "write_report_xlsx",
]
