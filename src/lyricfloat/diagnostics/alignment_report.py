"""L3 alignment statistics (logs/alignment.jsonl) as a workbook, to tune the thresholds
(SPEC v1.22). The log holds numbers only: no titles, no lyrics.
"""

from __future__ import annotations

import json
import logging
from pathlib import Path
from typing import Any

from lyricfloat.align.recorder import DROP_REASONS
from lyricfloat.lyrics.cache import default_cache_dir

log = logging.getLogger(__name__)

HELP = [
    "# 自動時間軸校正（L3）統計",
    "每一列是一首完整播完、送去對齊的歌。資料來自 %LOCALAPPDATA%\\LyricFloat\\logs\\alignment.jsonl，只有數字，不含歌名或歌詞。",
    "# 欄位",
    "採用：有把握的行達 6 成以上，這首改用對齊後的時間軸（SPEC v1.22）。",
    "有把握的行：與原時間差 3 秒內、且辨識信心 ≥ 0.2 的行；判斷範圍只算錄音涵蓋到的行。",
    "與原時間差中位數、0.3 秒內比例：對齊出的行開始時間和原始 LRC 的差距。LRC 本身也可能不準，所以這是兩者一致的程度，不是絕對誤差。",
    "平均信心：每行逐字辨識機率的平均，再對整首平均；用來調整 0.2 的門檻。",
    "耗時：人聲分離、對齊與整體（含載入模型）秒數。",
    "錄音起點：第一個錄到的取樣在歌曲裡的位置，錄音比歌曲早開始時是負數。起點修正：用整段錄音期間的進度回報取中位數，跟只用開錄當下那一筆相比差多少（v1.28 起）；換歌當下 Spotify 會提早回報新歌已開始，這裡常見 −0.3～−0.7 秒。",
    "# 沒錄成",
    "完整播放才會校正：中途暫停、拖動、開頭超過 15 秒才錄到、沒錄到結尾、靜音的歌都不處理，下次完整播放時再試。這一頁列出每次沒錄成的原因。",
]


def alignment_log_files(folder: Path | None = None) -> list[Path]:
    folder = folder or default_cache_dir() / "logs"
    files = [folder / "alignment.jsonl.2", folder / "alignment.jsonl.1", folder / "alignment.jsonl"]
    return [f for f in files if f.exists()]


def read_records(files: list[Path]) -> list[dict[str, Any]]:
    records = []
    for path in files:
        for line in path.read_text(encoding="utf-8").splitlines():
            try:
                record = json.loads(line)
            except ValueError:
                continue
            if isinstance(record, dict):
                records.append(record)
    return records


def write_alignment_xlsx(path: Path, records: list[dict[str, Any]]) -> None:
    from openpyxl import Workbook

    from lyricfloat.diagnostics.xlsx import Column, write_stats, write_table, write_text_sheet

    columns = [
        Column("at", "時間", 19),
        Column("language", "語言", 7),
        Column("ok", "成功", 7),
        Column("usable", "採用", 7, colors={"True": "D9EAD3", "False": "F4CCCC"}),
        Column("lines", "行數", 7),
        Column("judged", "判斷行數", 9),
        Column("sure", "有把握的行", 10),
        Column("sure_share", "有把握比例", 10),
        Column("median_abs_diff_s", "與原時間差中位數 (s)", 13),
        Column("within_0.3s", "0.3 秒內比例", 11),
        Column("median_shift_s", "整體平移 (s)", 11),
        Column("mean_confidence", "平均信心", 9),
        Column("word_lines", "有逐字的行", 10),
        Column("seconds", "錄音長度 (s)", 11),
        Column("pos0_s", "錄音起點 (s)", 10),
        Column("start_fix_s", "起點修正 (s)", 10),
        Column("reports", "進度回報數", 10),
        Column("separate_s", "人聲分離 (s)", 11),
        Column("align_s", "對齊 (s)", 9),
        Column("wall_s", "整體 (s)", 9),
    ]
    drops = [r for r in records if r.get("event") == "dropped"]
    aligned = [r for r in records if r.get("event") != "dropped"]
    rows = [{**r, "ok": str(r.get("ok")), "usable": str(r.get("usable")) if "usable" in r else None} for r in aligned]
    wb = Workbook()
    summary = wb.active
    summary.title = "摘要"
    write_table(wb.create_sheet("明細"), columns, rows, freeze="C2")
    drop_rows = [{"at": r.get("at"), "reason": DROP_REASONS.get(r.get("reason", ""), r.get("reason", ""))} for r in drops]
    write_table(wb.create_sheet("沒錄成"), [Column("at", "時間", 19), Column("reason", "原因", 22)], drop_rows)
    last = max(len(rows) + 1, 2)
    drop_last = max(len(drop_rows) + 1, 2)

    from openpyxl.utils import get_column_letter

    letters = {col.key: get_column_letter(i) for i, col in enumerate(columns, 1)}

    def rng(key: str) -> str:
        col = letters[key]
        return f"明細!{col}2:{col}{last}"

    write_stats(
        summary,
        1,
        [
            ("送去對齊的歌", f"=COUNTA({rng('at')})"),
            ("成功", f'=COUNTIFS({rng("ok")},"True")'),
            ("採用", f'=COUNTIFS({rng("usable")},"True")'),
            ("採用率（佔成功）", '=IF(B2=0,"—",B3/B2)'),
            ("華語：成功", f'=COUNTIFS({rng("language")},"zh",{rng("ok")},"True")'),
            ("華語：採用", f'=COUNTIFS({rng("language")},"zh",{rng("usable")},"True")'),
            ("華語：0.3 秒內比例（平均）", f'=IFERROR(AVERAGEIFS({rng("within_0.3s")},{rng("language")},"zh",{rng("ok")},"True"),"—")'),
            ("英文：成功", f'=COUNTIFS({rng("language")},"en",{rng("ok")},"True")'),
            ("英文：採用", f'=COUNTIFS({rng("language")},"en",{rng("usable")},"True")'),
            ("英文：0.3 秒內比例（平均）", f'=IFERROR(AVERAGEIFS({rng("within_0.3s")},{rng("language")},"en",{rng("ok")},"True"),"—")'),
            ("平均信心（全部成功的歌）", f'=IFERROR(AVERAGEIFS({rng("mean_confidence")},{rng("ok")},"True"),"—")'),
            ("整體耗時中位數 (s)", f'=IFERROR(MEDIAN({rng("wall_s")}),"—")'),
            ("沒錄成的播放", f"=COUNTA(沒錄成!A2:A{drop_last})"),
            *[(f"　{label}", f'=COUNTIFS(沒錄成!B2:B{drop_last},"{label}")') for label in dict.fromkeys(DROP_REASONS.values())],
        ],
    )
    write_text_sheet(wb.create_sheet("說明"), HELP)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
